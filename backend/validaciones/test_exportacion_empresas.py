"""
Pruebas de la exportación de empresas a Excel
(api/v1/servicios/exportacion_empresas.py) y de la cola de trabajos
(api/v1/servicios/cola_trabajos.py).

No necesitan la base ni levantar el server: el .xlsx se arma a partir de
filas escritas a mano y la cola se prueba con funciones comunes.
"""
import threading
import time
from datetime import timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from api.v1.servicios.busqueda_empresas import FiltrosEmpresa
from api.v1.servicios.cola_trabajos import ColaLlena, ColaTrabajos, EstadoTrabajo
from api.v1.servicios.exportacion_empresas import ENCABEZADOS, armar_xlsx

DATOS_COMPLETOS = {
    "nombre_empresa": "Sabores Serranos",
    "rubro": "AGROALIMENTO",
    "ubicacion": {
        "zona": "Zona 5", "departamento": "CALAMUCHITA", "localidad": "Villa General Belgrano",
        "calle": "San Martín", "numero": 123, "latitud": -31.97, "longitud": -64.55,
    },
    "descripcion": "Dulces artesanales",
    "correo_empresa": "contacto@sabores.com",
    "redes": {"numero_contacto": "+5493511234567", "instagram": "@sabores"},
}
# (id_empresa, cuit, razon_social, datos_publico, estado, motivo_rechazo)
FILA_COMPLETA = (7, 30712345678, "Sabores Serranos SRL", DATOS_COMPLETOS, 1, None)
FILA_INCOMPLETA = (1, 20111111112, "Carga Manual SA", {"rubro": "AGTECH", "ubicacion": {"zona": "Zona 1"}}, 2, "Faltan datos")


def _abrir(contenido: bytes):
    return load_workbook(BytesIO(contenido))


def _esperar(cola: ColaTrabajos, trabajo, segundos: float = 3.0):
    limite = time.monotonic() + segundos
    while trabajo.terminado_en is None and time.monotonic() < limite:
        time.sleep(0.01)
    return trabajo


# --- Excel ---------------------------------------------------------------


def test_xlsx_tiene_encabezados_y_una_fila_por_empresa():
    hoja = _abrir(armar_xlsx([FILA_COMPLETA, FILA_INCOMPLETA], FiltrosEmpresa()))["Empresas"]
    assert [celda.value for celda in hoja[1]] == ENCABEZADOS
    assert hoja.max_row == 3


def test_xlsx_aplana_el_json_en_columnas():
    hoja = _abrir(armar_xlsx([FILA_COMPLETA], FiltrosEmpresa()))["Empresas"]
    fila = dict(zip(ENCABEZADOS, (celda.value for celda in hoja[2])))
    assert fila["Nombre"] == "Sabores Serranos"
    assert fila["Departamento"] == "CALAMUCHITA"
    assert fila["Instagram"] == "@sabores"
    assert fila["Estado"] == "Aprobada"
    assert fila["CUIT"] == "30712345678"  # texto, no número


def test_xlsx_exporta_empresa_incompleta_con_celdas_vacias():
    hoja = _abrir(armar_xlsx([FILA_INCOMPLETA], FiltrosEmpresa()))["Empresas"]
    fila = dict(zip(ENCABEZADOS, (celda.value for celda in hoja[2])))
    assert fila["Razón social"] == "Carga Manual SA"
    assert fila["Nombre"] is None and fila["Correo"] is None
    assert fila["Estado"] == "Rechazada" and fila["Motivo de rechazo"] == "Faltan datos"


def test_xlsx_no_convierte_texto_del_usuario_en_formula():
    datos = {**DATOS_COMPLETOS, "descripcion": "=1+1"}
    hoja = _abrir(armar_xlsx([(7, 30712345678, "X SA", datos, 1, None)], FiltrosEmpresa()))["Empresas"]
    celda = hoja.cell(row=2, column=ENCABEZADOS.index("Descripción") + 1)
    assert celda.value == "=1+1"
    assert celda.data_type == "s"


def test_xlsx_incluye_los_filtros_usados():
    libro = _abrir(armar_xlsx([FILA_COMPLETA], FiltrosEmpresa(rubro="AGROALIMENTO", estado=1)))
    resumen = {fila[0].value: fila[1].value for fila in libro["Filtros"].iter_rows()}
    assert resumen["Empresas exportadas"] == 1
    assert resumen["Filtro: rubro"] == "AGROALIMENTO"
    assert resumen["Filtro: estado"] == "Aprobada"
    assert resumen["Filtro: zona"] == "(sin filtrar)"


# --- Cola ----------------------------------------------------------------


def test_cola_ejecuta_el_trabajo_y_guarda_el_resultado():
    cola = ColaTrabajos("test")
    trabajo = _esperar(cola, cola.encolar(lambda: b"archivo", solicitante=1))
    assert trabajo.estado == EstadoTrabajo.LISTO
    assert trabajo.resultado == b"archivo"
    assert cola.obtener(trabajo.id) is trabajo


def test_cola_marca_error_sin_filtrar_el_detalle():
    def falla():
        raise RuntimeError("ORA-00942: detalle interno")

    cola = ColaTrabajos("test")
    trabajo = _esperar(cola, cola.encolar(falla, solicitante=1))
    assert trabajo.estado == EstadoTrabajo.ERROR
    assert "ORA-" not in trabajo.error


def test_cola_procesa_de_a_uno_y_en_orden():
    cola = ColaTrabajos("test")
    soltar = threading.Event()
    orden = []

    primero = cola.encolar(lambda: (soltar.wait(3), orden.append("a")), solicitante=1)
    segundo = cola.encolar(lambda: orden.append("b"), solicitante=1)
    tercero = cola.encolar(lambda: orden.append("c"), solicitante=1)
    time.sleep(0.05)

    assert primero.estado == EstadoTrabajo.PROCESANDO
    assert segundo.estado == tercero.estado == EstadoTrabajo.EN_COLA
    assert (cola.posicion(primero), cola.posicion(segundo), cola.posicion(tercero)) == (0, 0, 1)

    soltar.set()
    _esperar(cola, tercero)
    assert orden == ["a", "b", "c"]
    # un trabajo que falla no frena a los siguientes
    cuarto = _esperar(cola, cola.encolar(lambda: 1 / 0, solicitante=1))
    quinto = _esperar(cola, cola.encolar(lambda: "ok", solicitante=1))
    assert (cuarto.estado, quinto.estado) == (EstadoTrabajo.ERROR, EstadoTrabajo.LISTO)


def test_cola_llena_rechaza_trabajos_nuevos():
    cola = ColaTrabajos("test", maximo_en_espera=1)
    soltar = threading.Event()
    cola.encolar(soltar.wait, solicitante=1)  # lo toma el trabajador
    time.sleep(0.05)
    cola.encolar(lambda: None, solicitante=1)  # queda esperando: cola llena
    with pytest.raises(ColaLlena):
        cola.encolar(lambda: None, solicitante=1)
    soltar.set()


def test_cola_borra_los_resultados_vencidos():
    cola = ColaTrabajos("test", ttl_resultado=timedelta(seconds=0))
    trabajo = _esperar(cola, cola.encolar(lambda: b"x", solicitante=1))
    time.sleep(0.01)
    assert cola.obtener(trabajo.id) is None
