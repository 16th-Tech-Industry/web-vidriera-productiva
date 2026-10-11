# backend/api/v1/servicios/exportacion_empresas.py
"""
Exportación de empresas a Excel (.xlsx).

Respeta los mismos filtros que el buscador del dashboard admin: usa
buscar_empresas() (api/v1/servicios/busqueda_empresas.py) sin paginar.

Dos piezas:
- armar_xlsx(): filas + filtros -> bytes del archivo. No toca la base, así
  que se puede testear sola (ver validaciones/test_exportacion_empresas.py).
- generar_xlsx_empresas(): consulta + armar_xlsx(). Es lo que se encola
  (ver endpoints/dashboard/empresas_exportacion.py).
"""
import json
from dataclasses import asdict
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from api.v1.schemas.empresas import ESTADO_APROBADA, ESTADO_PENDIENTE, ESTADO_RECHAZADA
from api.v1.servicios.busqueda_empresas import FiltrosEmpresa, buscar_empresas

MEDIA_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

COLUMNAS_EXPORTACION = "id_empresa, cuit, razon_social, datos_publico, estado, motivo_rechazo"

NOMBRES_ESTADO = {
    ESTADO_PENDIENTE: "Pendiente",
    ESTADO_APROBADA: "Aprobada",
    ESTADO_RECHAZADA: "Rechazada",
}

ENCABEZADOS = [
    "ID", "CUIT", "Razón social", "Nombre", "Rubro", "Zona", "Departamento", "Localidad",
    "Calle", "Número", "Latitud", "Longitud", "Descripción", "Correo", "Teléfono",
    "Facebook", "Instagram", "Página web", "Estado", "Motivo de rechazo",
]
ANCHO_MINIMO = 10
ANCHO_MAXIMO = 50


def _fila_a_registro(fila) -> list:
    """Aplana una fila de `empresas` (con su JSON) al orden de ENCABEZADOS.

    A diferencia de los listados de la API, acá una empresa con el JSON
    incompleto NO se saltea: se exporta con las celdas que falten vacías,
    para que el admin pueda detectarla en la planilla.
    """
    id_empresa, cuit, razon_social, datos_publico, estado, motivo_rechazo = fila
    datos = datos_publico if isinstance(datos_publico, dict) else json.loads(datos_publico or "{}")
    ubicacion = datos.get("ubicacion") or {}
    redes = datos.get("redes") or {}
    return [
        int(id_empresa),
        str(int(cuit)),  # como texto: Excel mostraría un número de 11 dígitos como 2,03E+10
        razon_social,
        datos.get("nombre_empresa"),
        datos.get("rubro"),
        ubicacion.get("zona"),
        ubicacion.get("departamento"),
        ubicacion.get("localidad"),
        ubicacion.get("calle"),
        ubicacion.get("numero"),
        ubicacion.get("latitud"),
        ubicacion.get("longitud"),
        datos.get("descripcion"),
        datos.get("correo_empresa"),
        redes.get("numero_contacto"),
        redes.get("facebook"),
        redes.get("instagram"),
        redes.get("pagina_web"),
        NOMBRES_ESTADO.get(int(estado), str(estado)),
        motivo_rechazo,
    ]


def _escribir_fila(hoja, valores: list) -> None:
    hoja.append(valores)
    # openpyxl interpreta como FÓRMULA cualquier texto que empiece con "=".
    # Estos datos los cargan los usuarios: se fuerzan a texto para que una
    # descripción como "=HYPERLINK(...)" no se ejecute al abrir la planilla.
    for celda in hoja[hoja.max_row]:
        if isinstance(celda.value, str):
            celda.data_type = "s"


def _ajustar_anchos(hoja) -> None:
    for indice, columna in enumerate(hoja.columns, start=1):
        largo = max((len(str(celda.value)) for celda in columna if celda.value is not None), default=0)
        hoja.column_dimensions[get_column_letter(indice)].width = max(ANCHO_MINIMO, min(largo + 2, ANCHO_MAXIMO))


def armar_xlsx(filas: list, filtros: FiltrosEmpresa) -> bytes:
    """Arma el libro: hoja "Empresas" con los datos y hoja "Filtros" con los
    criterios que se usaron, para que la planilla se explique sola."""
    libro = Workbook()

    hoja = libro.active
    hoja.title = "Empresas"
    hoja.append(ENCABEZADOS)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in filas:
        _escribir_fila(hoja, _fila_a_registro(fila))
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    _ajustar_anchos(hoja)

    resumen = libro.create_sheet("Filtros")
    _escribir_fila(resumen, ["Generado", datetime.now().strftime("%d/%m/%Y %H:%M")])
    _escribir_fila(resumen, ["Empresas exportadas", len(filas)])
    for nombre, valor in asdict(filtros).items():
        if nombre == "estado" and valor is not None:
            valor = NOMBRES_ESTADO.get(valor, valor)
        _escribir_fila(resumen, [f"Filtro: {nombre}", "(sin filtrar)" if valor is None else valor])
    for celda in resumen["A"]:
        celda.font = Font(bold=True)
    _ajustar_anchos(resumen)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_xlsx_empresas(filtros: FiltrosEmpresa) -> bytes:
    """Trae todas las empresas que cumplen los filtros y devuelve el .xlsx."""
    filas, _ = buscar_empresas(COLUMNAS_EXPORTACION, filtros)
    return armar_xlsx(filas, filtros)
